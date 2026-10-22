import os
import time
import threading
import logging
from datetime import datetime, timedelta, date
from functools import wraps

from requests.exceptions import ConnectionError, ReadTimeout
from openpyxl import load_workbook
import telebot
from telebot import types
import re

import UpdateFileShedule

# Настройка логирования
logging.basicConfig(
   format='%(asctime)s - %(levelname)s - %(message)s',
   level=logging.INFO,
   filename=os.path.join('/root/Flash/TgLog/', 'bot.log'), 
   filemode='a'
)
logger = logging.getLogger(__name__)

bot = telebot.TeleBot("7703428815:AAFCJFu9o3JTuYgFb6KgegfiTi9HMlns1Yw")
ALLOWED_USER_IDS = [942120524, 1219623875]        # Список пользователей тг бота
LOCAL_FILE_PATH = "/root/Flash/Расписание.xlsx"   # Путь к файлу расписания

CountStr = 10   # Количество строк, которые пойдут в ответ в бот

days_mapping = { 1: 'пн', 2: 'вт', 3: 'ср', 4: 'чт', 5: 'пт', 6: 'сб', 7: 'вс' }   # словарь дней

CurrentWeekColor = 'FFC6EFCE'   # Цвет текущей недели
ColorNoStudy     = '6'          # Цвет для занятия которое отменилось

def openShedule():
   try:
      if not os.path.exists(LOCAL_FILE_PATH):
         logger.error(f"Файл не найден: {LOCAL_FILE_PATH}")
         return None

      wb = load_workbook(LOCAL_FILE_PATH)
      ws = wb.active
      
      return ws

   except Exception as e:
      logger.error(f"Ошибка при чтении локального файла: {e}")
      return None

# Найти расписание на указанный день недели
def find_shedule(day, IsNeedNextMonday):
   file = openShedule()
   if file is None: return "❌ Не удалось загрузить расписание"

   # Ищем колонку с сегодняшним днем в первой строке
   target_column_idx = None
   for col_idx in range(1, file.max_column + 1):
      cell = file.cell(row=1, column=col_idx)
      cell_value = str(cell.value) if cell.value else ""
      if day in cell_value.lower():
         target_column_idx = col_idx
         break
   
   if target_column_idx is None: return f"❌ Колонка для дня '{day}' не найдена"

   target_row_idx = None
   target_cell_value = None
   
   # Идем снизу вверх (от последней строки к первой)
   for row_idx in range(file.max_row, 0, -1):
      cell = file.cell(row=row_idx, column=target_column_idx)
      cell_value = str(cell.value) if cell.value else ""
      if day in cell_value.lower():
         if IsNeedNextMonday:
            target_row_idx = row_idx
            target_cell_value = cell_value
            break
         else:
            if get_cell_color(file, row_idx, target_column_idx) == CurrentWeekColor:
               target_row_idx = row_idx
               target_cell_value = cell_value
               break

   if target_row_idx is None:
      logger.warning(f"Не найдена строка с днем '{day}'")
      return f"❌ Не найдена строка с днем '{day}'"

   schedule_text = ""
   current_time = None
   
   '''current_weekday = datetime.now().isoweekday()
   current_day_name = days_mapping[current_weekday]

   if day == current_day_name:
      if current_day_name == 'вс': current_time = datetime.now()
      else:                        current_time = datetime.now()
      #else:                        current_time = datetime.now() + timedelta(days=1)
   else:
      current_time = datetime.now()
      
   #if current_time != None: schedule_text += f"{day} {current_time.day:02d}.{current_time.month:02d}\n\n"
   
   if IsNeedNextMonday: processed_cell_value = target_cell_value
   else:
      date_str = f"{current_time.day:02d}.{current_time.month:02d}"
      date_pattern = r'\b\d{2}\.\d{2}\b'
      if re.search(date_pattern, target_cell_value): processed_cell_value = re.sub(date_pattern, date_str, target_cell_value)
      else:                                          processed_cell_value = target_cell_value
'''
   processed_cell_value = target_cell_value
   schedule_text += f"{processed_cell_value}\n\n"

   rows_added = 0
   for row_idx in range(target_row_idx + 1, file.max_row + 1):
      if rows_added >= CountStr:
         break
      
      cell = file.cell(row=row_idx, column=target_column_idx)
      cell_value = str(cell.value) if cell.value else ""
      
      # Пропускаем пустые ячейки
      if cell_value and cell_value != 'None' and cell_value != 'nan':
         if get_cell_color(file, row_idx, target_column_idx) == ColorNoStudy: schedule_text += f"❌ {cell_value} - отмена\n"   # если помечено цветом отмены, то пишу что отмена
         else:                                                                schedule_text += f"✏️ {cell_value}\n"
         schedule_text += f"------------------------------------\n"

      rows_added += 1

   if rows_added == 0:
      logger.warning(f"Не найдено расписания для {day}")
      schedule_text = f"❌ На сегодня расписания нет"
   
   return schedule_text

def get_next_monday_schedule(chatId):
   day_index = datetime.now().isoweekday()
   day   = days_mapping[day_index]

   if (day == 'сб'): 
      bot.send_message(chatId, "Выгружаю", reply_markup=create_main_menu(True)) 
      return find_shedule('пн', True)
   
   elif (day == 'вс'): 
      bot.send_message(chatId, "Выгружаю", reply_markup=create_main_menu1())
      return find_shedule('пн', True)
   
   else:                            
      bot.send_message(chatId, "Обновил меню", reply_markup=create_main_menu(False)) 
      return "Нет смысла смотреть, сегодня не суббота/воскресенье"

def get_today_schedule(chatId):
   day_index = datetime.now().isoweekday()
   day   = days_mapping[day_index]

   if (day == 'сб'):   bot.send_message(chatId, "Выгружаю", reply_markup=create_main_menu(True))
   elif (day == 'вс'): bot.send_message(chatId, "Выгружаю", reply_markup=create_main_menu1())  
   else:               bot.send_message(chatId, "Выгружаю", reply_markup=create_main_menu(False)) 

   return find_shedule(day, False)


def get_tomorrow_schedule(chatId):
   day_index = (datetime.now().isoweekday() + 1) % 8  # Завтрашний день
   
   if day_index == 0: day_index += 1

   day = days_mapping[day_index]

   day_str = days_mapping[datetime.now().isoweekday()]

   if (day_str == 'сб'):   bot.send_message(chatId, "Выгружаю", reply_markup=create_main_menu(True))
   elif (day_str == 'вс'): bot.send_message(chatId, "Выгружаю", reply_markup=create_main_menu1())  
   else:                   bot.send_message(chatId, "Выгружаю", reply_markup=create_main_menu(False)) 

   return find_shedule(day, False)

def create_main_menu1():
   markup = types.ReplyKeyboardMarkup(resize_keyboard=True, row_width=2)
   btn1 = types.KeyboardButton('📝 Расписание на сегодня')
   btn3 = types.KeyboardButton('📝 Расписание на понедельник')

   markup.row(btn1, btn3)  
                    
   return markup


def create_main_menu(IsNeedNextMonday):
   markup = types.ReplyKeyboardMarkup(resize_keyboard=True, row_width=2)
   btn1 = types.KeyboardButton('📝 Расписание на сегодня')
   btn2 = types.KeyboardButton('📝 Расписание на завтра')

   if(IsNeedNextMonday):
      btn3 = types.KeyboardButton('📝 Расписание на понедельник')
      markup.row(btn1, btn2)  
      markup.add(btn3)      
   else : markup.row(btn1, btn2)                    

   return markup

# Декоратор для проверки пользователя
def allowed_users_only(func):
   @wraps(func)
   def wrapper(message, *args, **kwargs):
      if message.from_user.id not in ALLOWED_USER_IDS:
         bot.send_message(message.chat.id, "Доступ запрещен!")
         logger.info(f"Пользователь {message.from_user.first_name}/{message.from_user.id} ломится")
         return
      return func(message, *args, **kwargs)
   return wrapper

def get_cell_color(file, row, column):
   try:
      cell = file.cell(row=row, column=column)
      color_value = cell.fill.start_color.index
              
      if isinstance(color_value, int): return str(color_value)
      else:                            return color_value  
   except Exception as e:
      logger.error(f"Ошибка при получении цвета ячейки ({row}, {column}): {e}")
      return None

# Дополнительные обработчики с проверкой
@bot.message_handler(commands=['start'])
@allowed_users_only
def handle_start(message: types.Message):
   logger.info(f"Пользователь {message.from_user.first_name}/{message.from_user.id} авторизован")
   bot.send_message(message.chat.id, f"Привет, {message.from_user.first_name}", reply_markup=create_main_menu())
   day_index = datetime.now().isoweekday()
   day = days_mapping[day_index]

   if (day == 'сб' or day == 'вс'): bot.send_message(message.chat.id, f"Привет, {message.from_user.first_name}", reply_markup=create_main_menu(True)) 
   else:                            bot.send_message(message.chat.id, f"Привет, {message.from_user.first_name}", reply_markup=create_main_menu(False)) 

# Обработчик текстовых сообщений (кнопки меню)
@bot.message_handler(content_types=['text'], func=lambda message: not message.text.startswith('/'))
@allowed_users_only
def handle_menu_buttons(message: types.Message):
   if message.text == '📝 Расписание на сегодня':
      schedule = get_today_schedule(message.chat.id)
      bot.send_message(message.chat.id, schedule)

   elif message.text == '📝 Расписание на завтра':
      schedule = get_tomorrow_schedule(message.chat.id)
      bot.send_message(message.chat.id, schedule)

   elif message.text == '📝 Расписание на понедельник':
      schedule = get_next_monday_schedule(message.chat.id)
      bot.send_message(message.chat.id, schedule)

   else:
      bot.send_message(message.chat.id, "Не понимаю команду", reply_markup=create_main_menu())

@bot.message_handler(commands=['log'])
@allowed_users_only
def handle_log(message: types.Message):
   try:
      today = date.today()
      today_str = today.strftime('%Y-%m-%d')
      
      log_file_path = os.path.join('/root/Flash/TgLog/', 'bot.log')
      
      if not os.path.exists(log_file_path):
         bot.send_message(message.chat.id, "Лог-файл не найден!")
         return

      today_logs = []
      with open(log_file_path, 'r', encoding='utf-8') as file:
         for line in file:
            if line.startswith(today_str):
               today_logs.append(line.strip())
      
      if not today_logs:
         bot.send_message(message.chat.id, f"Логи за сегодня ({today_str}) не найдены!")
         return
      
      log_text = "\n".join(today_logs)
      
      if len(log_text) > 4000:
         for i in range(0, len(log_text), 4000):
            chunk = log_text[i:i+4000]
            bot.send_message(message.chat.id, f"<code>{chunk}</code>", parse_mode='HTML')
      else: bot.send_message(message.chat.id, f"<code>{log_text}</code>", parse_mode='HTML')
         
   except Exception as e: bot.send_message(message.chat.id, f"Ошибка при чтении логов: {str(e)}")

def update_file():

   success = UpdateFileShedule.download_file()
   
   if    success: pass
   else: logger.warning("Не удалось обновить файл расписания")
   
   return success

def UpdateFileLoop():
   logger.info("Запустил задачу по обновлению файла расписания")
   
   while True: 
      try:
         while True:
            update_file()
            time.sleep(7200)  # Спим 2 часа
               
      except Exception as e:
         logger.error(f"ПРОЦЕСС УПАЛ: {e} - перезапуск через 10 сек")
         time.sleep(10)

def start_bot():
   logger.info("Бот запущен")
   while True:
      try:
         bot.polling(none_stop=True, interval=2, timeout=60, long_polling_timeout=60)
      except ConnectionError:
         logger.warning("Ошибка соединения")
         time.sleep(10)
      except ReadTimeout:
         logger.warning("Таймаут соединения")
         time.sleep(10)
      except Exception as e:
         logger.error(f"Unexpected error: {e}")
         logger.info("Перезапуск")
         time.sleep(10)

if __name__ == '__main__':
   FileUpdateThread = threading.Thread(target=UpdateFileLoop, daemon=True)
   FileUpdateThread.start()

   start_bot()
